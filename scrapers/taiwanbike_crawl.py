# -*- coding: utf-8 -*-
"""Scrape TBA (中華民國自行車協會) ultra-distance certification results: the
taiwanbike.org listing page links to public Google Sheets (2022+); fetch each as
xlsx and parse every tab with the tolerant mapper. Finish time only; rank_overall
None (frontend re-ranks); sub-1h placeholders dropped.
Output: data/processed/taiwanbike_2022_2026.json
"""
import io
import json
import os
import re
import sys
from datetime import datetime, timezone

sys.path.insert(0, os.path.dirname(__file__))
import common  # noqa: E402
import taiwanbike_parse as tp  # noqa: E402

LISTING = "https://taiwanbike.org/index.php/2009-07-18-17-11-16"
SHEET_XLSX = "https://docs.google.com/spreadsheets/d/{}/export?format=xlsx"
MIN_FINISH_SEC = 3600
OUT = os.path.join(os.path.dirname(__file__), "..", "data", "processed",
                   "taiwanbike_2022_2026.json")
_YEAR = re.compile(r"^20\d{2}\s*")


def records_from_sheet(tabs, title, year, scraped_at):
    """tabs: dict{tab_name: list[list]} (first row = header). Reads the per-rider
    roster tab(s); skips obvious rank-only re-sort tabs to avoid double counting."""
    race = _YEAR.sub("", title).strip()
    out, seen = [], set()
    for name, rows in tabs.items():
        if any(k in name for k in ("排序", "排名")) and "總表" not in name:
            continue  # rank-sort tabs duplicate the roster
        if not rows or len(rows) < 2:
            continue
        header = rows[0]
        for row in rows[1:]:
            m = tp.map_row(header, row)
            if not m:
                continue
            fs = common.time_to_seconds(m["finish_time"])
            if fs is None or fs < MIN_FINISH_SEC:
                continue
            key = (m["bib"], m["name"], int(fs))
            if key in seen:
                continue
            seen.add(key)
            out.append(common.make_record(
                source_platform="taiwanbike.org",
                source_url=LISTING, source_format="xlsx",
                race_name_raw=f"{race} {year}", year=year,
                result_label=race, category_raw=m["category"],
                gender=m["gender"], age_group=None,
                rank_overall=None, bib=m["bib"], team=m["team"],
                name_raw=m["name"], finish_time=m["finish_time"],
                scraped_at=scraped_at))
    return out


def _xlsx_tabs(content):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    tabs = {}
    for ws in wb.worksheets:
        tabs[ws.title] = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return tabs


def crawl(session=None):
    session = session or common.make_session()
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    html = common.polite_get(session, LISTING).text
    links = tp.extract_sheet_links(html)
    print(f"  {len(links)} sheet links (2022+)")
    out = []
    for lk in links:
        try:
            content = common.polite_get(session, SHEET_XLSX.format(lk["sheet_id"])).content
            tabs = _xlsx_tabs(content)
            recs = records_from_sheet(tabs, lk["title"], lk["year"], now)
            print(f"    {lk['title']}: {len(recs)} rows")
            out.extend(recs)
        except Exception as e:  # one bad sheet must not abort the run
            print(f"    ! skip {lk['title']}: {e}")
    return out


def main():
    recs = crawl()
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(recs, f, ensure_ascii=False, indent=1)
    print(f"\n  {len(recs)} rows -> {os.path.relpath(OUT)}")


if __name__ == "__main__":
    main()
